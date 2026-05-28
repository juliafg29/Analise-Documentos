# PaddleOCR

# Ainda não cheguei a um acordo sobre qual código está apresentando melhores resultados.

# Claude 4.0 + Codex


"""
Pipeline OCR para documentos brasileiros.

Etapas:
1. OCR com PaddleOCR
2. Agrupamento espacial em linhas
3. Normalização e correções comuns de OCR
4. Identificação de rótulos
5. Busca de valores:
   - mesma linha: "REGISTRO GERAL 28.XXX.102-5"
   - mesma linha com vários rótulos:
     "REGISTRO GERAL 28.XXX.102-5 DATA DE EXPEDICAO 19/04/2024"
   - linha abaixo: "DATA DE NASCIMENTO" / "10/05/1990"
6. Classificação estruturada dos campos
7. Geração de XML
"""

import os
import re
import json
import unicodedata
import xml.etree.ElementTree as ET

from pathlib import Path
from datetime import datetime
from dataclasses import dataclass, asdict
from difflib import SequenceMatcher

import cv2
import numpy as np
from paddleocr import PaddleOCR


# ============================================================
# CONFIGURAÇÕES
# ============================================================

MIN_SCORE = 0.50


EXTENSOES_IMAGEM = {
    ".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".webp"
}


# ============================================================
# NORMALIZAÇÃO
# ============================================================

def normalizar_texto(texto: str) -> str:
    if not texto:
        return ""

    nfkd = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in nfkd if not unicodedata.combining(c))
    texto = texto.upper()

    texto = texto.replace("|", " ")
    texto = texto.replace("_", " ")
    texto = re.sub(r"[^A-Z0-9\s\.\-\/:]", " ", texto)
    texto = re.sub(r"\s+", " ", texto)

    return texto.strip()


CORRECOES_OCR = {
    "DATADE": "DATA DE",
    "DATAEXPEDICAO": "DATA EXPEDICAO",
    "DATANASCIMENTO": "DATA NASCIMENTO",
    "DTNASC": "DT NASC",
    "DTEXPEDICAO": "DT EXPEDICAO",
    "EXPEDICA0": "EXPEDICAO",
    "NASCIMENT0": "NASCIMENTO",
    "EMISSA0": "EMISSAO",
    "ORG.AO": "ORGAO",
    "ORGA0": "ORGAO",
    "CPF/MF": "CPF",
}


def aplicar_correcoes_ocr(texto: str) -> str:
    texto = normalizar_texto(texto)

    for errado, certo in CORRECOES_OCR.items():
        texto = texto.replace(errado, certo)

    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


def similaridade(a: str, b: str) -> float:
    return SequenceMatcher(None, aplicar_correcoes_ocr(a), aplicar_correcoes_ocr(b)).ratio()


# ============================================================
# REGEX
# ============================================================

REGEX_CPF = re.compile(
    r"\b\d{3}[\.\s]?\d{3}[\.\s]?\d{3}[-\s]?\d{2}\b"
)

#REGEX_RG = re.compile(
#    r"\b\d{1,2}[\.\s]?\d{3}[\.\s]?\d{3}[-\s]?[0-9Xx]\b"
#)

REGEX_DATA = re.compile(
    r"\b\d{1,2}[\/\-.](?:\d{1,2}|[A-Z]{3})[\/\-.]\d{2,4}\b",
    re.IGNORECASE
)

REGEX_LOCAL_UF = re.compile(
    r"\b[A-Z][A-Z\s]{2,}\s[-\/]?\s[A-Z]{2}\b"
)


# ============================================================
# RÓTULOS
# ============================================================

ROTULOS_CAMPOS = {
    "nome": [
        "NOME",
        "NOME COMPLETO",
        "NOME CIVIL",
        "NOME SOCIAL",
    ],
    "cpf": [
        "CPF",
        "C.P.F",
        "C P F",
        "CADASTRO DE PESSOA FISICA",
        "CADASTRO DE PESSOAS FISICAS",
    ],
    "rg": [
        "REGISTRO GERAL",
        "R.G",
        "RG",
        "GERAL",
        #"IDENTIDADE",
        "DOC IDENTIDADE",
       # "DOCUMENTO DE IDENTIDADE",
    ],
    "data_nascimento": [
        "DATA DE NASCIMENTO",
        "DATA NASCIMENTO",
        "DT NASCIMENTO",
        "DATA NASC",
        "DT NASC",
        "NASCIMENTO",
        "NASC",
    ],
    "data_expedicao": [
        "DATA DE EXPEDICAO",
        "DT EXPEDICAO",
        "DATA EXPEDICAO",
        "EXPEDICAO",
        "DATA DE EMISSAO",
        "DT EMISSAO",
        "EMISSAO",
        "EXPED",
    ],
    "local_nascimento": [
        "NATURALIDADE",
        "LOCAL DE NASCIMENTO",
        "LOCAL NASCIMENTO",
        "LOCAL NASC",
        "NATURAL DE",
        "MUNICIPIO NASCIMENTO",
    ],
    "orgao_emissor": [
        "ORGAO EXPEDIDOR",
        "ORG EXPEDIDOR",
        "ORGAO EMISSOR",
        "ORGAO EMISSOR UF",
        "SSP",
        "DETRAN",
        "IFP",
    ],
 #   "filiacao": [
 #       "FILIACAO",
 #       "NOME DA MAE",
 #       "NOME DO PAI",
 #       "MAE",
 #       "PAI",
 #   ],
}

CAMPOS_DATA = {
    "data_nascimento",
    "data_expedicao",
}

PALAVRAS_RUIDO = {
    "REPUBLICA",
    "FEDERATIVA",
    "BRASIL",
    "VALIDA",
    "TODO",
    "TERRITORIO",
    "NACIONAL",
    "CARTEIRA",
    "IDENTIDADE",
    "SECRETARIA",
    "SEGURANCA",
    "PUBLICA",
    "DETRAN",
    "SENATRAN",
    "DENATRAN",
    "MINISTERIO",
    "GOVERNO",
}


# ============================================================
# LIMPEZA E VALIDAÇÃO
# ============================================================

def limpar_cpf(valor: str) -> str | None:
    digitos = re.sub(r"\D", "", valor or "")

    # CPF deve ter 11 dígitos
    if len(digitos) != 11:
        return None

    # Rejeita sequências iguais (000..., 111..., etc.)
    if len(set(digitos)) == 1:
        return None

    # Validação do 1º dígito
    soma = sum(int(digitos[i]) * (10 - i) for i in range(9))
    resto = (soma * 10) % 11
    digito_1 = 0 if resto == 10 else resto

    if digito_1 != int(digitos[9]):
        return None
    
    # Validação do 2º dígito
    soma = sum(int(digitos[i]) * (11 - i) for i in range(10))
    resto = (soma * 10) % 11
    digito_2 = 0 if resto == 10 else resto

    if digito_2 != int(digitos[10]):
        return None

    # Retorna CPF formatado
    return f"{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:]}"


def limpar_rg(valor: str) -> str | None:
    if not valor:
        return None

    valor_norm = aplicar_correcoes_ocr(valor)

    # Remove rótulos e palavras que podem vir junto
    valor = re.sub(r"\b(REGISTRO|GERAL|RG|R\.G|DOC|DOCUMENTO|IDENTIDADE)\b", " ", valor)
    valor = re.sub(r"\b(DATA|EXPEDICAO|NOME|CPF|NASCIMENTO|NATURALIDADE)\b.*", " ", valor)

    # Mantém letras, números, ponto, hífen e barra
    valor = re.sub(r"[^A-Z0-9.\-/]", " ", valor)
    valor = re.sub(r"\s+", " ", valor).strip()

    # Extrai candidatos com pelo menos 4 caracteres úteis
    candidatos = re.findall(r"[A-Z0-9][A-Z0-9.\-/]{3,}", valor)

    if not candidatos:
        return None

    # Preferir candidato com mais dígitos
    candidatos = sorted(
        candidatos,
        key=lambda x: sum(c.isdigit() for c in x),
        reverse=True
    )

    return candidatos[0]


def limpar_data(valor: str) -> str | None:
    m = REGEX_DATA.search(valor or "")

    if not m:
        return None

    partes = re.split(r"[\/\-.]", m.group())
    if len(partes) != 3:
        return None

    dia, mes, ano = partes

    #if len(ano) == 2:
    #    ano = "19" + ano if int(ano) > 30 else "20" + ano

    try:
        data = datetime(int(ano), int(mes), int(dia))
        return data.strftime("%d/%m/%Y")
    except ValueError:
        return None


def limpar_texto(valor: str) -> str | None:
    valor = aplicar_correcoes_ocr(valor)
    valor = re.sub(r"[^A-Z0-9\s\/\-.]", " ", valor)
    valor = re.sub(r"\s+", " ", valor).strip()

    return valor if len(valor) >= 2 else None


def parece_nome(valor: str) -> bool:
    valor = limpar_texto(valor)

    if not valor:
        return False

    palavras = valor.split()

    if len(palavras) < 2:
        return False

    if any(char.isdigit() for char in valor):
        return False

    if any(p in PALAVRAS_RUIDO for p in palavras):
        return False

    return True


def parece_local_nascimento(valor: str) -> bool:
    valor = limpar_texto(valor)

    if not valor:
        return False

    if any(char.isdigit() for char in valor):
        return False

    if any(p in PALAVRAS_RUIDO for p in valor.split()):
        return False

    return len(valor.split()) >= 1


def limpar_valor_por_campo(campo: str, valor: str) -> str | None:
    if campo == "cpf":
        return limpar_cpf(valor)

    if campo == "rg":
        return limpar_rg(valor)

    if campo in CAMPOS_DATA:
        return limpar_data(valor)

    if campo == "nome":
        valor = limpar_texto(valor)
        return valor if parece_nome(valor) else None

    if campo == "local_nascimento":
        valor = limpar_texto(valor)
        return valor if parece_local_nascimento(valor) else None

    return limpar_texto(valor)


# ============================================================
# ESTRUTURA DE LINHA OCR
# ============================================================

@dataclass
class LinhaOCR:
    indice: int
    texto: str
    norm: str
    score_medio: float = 0.0
    tokens: list[dict] | None = None
    x_min: int | None = None
    y_min: int | None = None
    x_max: int | None = None
    y_max: int | None = None
    cx: int | None = None
    cy: int | None = None

def calcular_confianca_etiqueta(
    campo: str,
    valor: str,
    score_ocr: float,
    score_rotulo: float = 0.0,
    origem: str = "",
) -> float:
    pesos_origem = {
        "mesma_linha": 1.00,
        "linha_abaixo": 0.90,
        "regex_global": 0.85,
        "regex_linha": 0.75,
        "fallback_unica_data_solta": 0.55,
        "regex_local_uf": 0.55,
    }

    peso_origem = pesos_origem.get(origem, 0.50)

    score_validacao = 0.70

    if campo == "cpf" and limpar_cpf(valor):
        score_validacao = 1.00
    elif campo in CAMPOS_DATA and limpar_data(valor):
        score_validacao = 0.95
    elif campo == "nome" and parece_nome(valor):
        score_validacao = 0.90
    elif campo == "rg" and limpar_rg(valor):
        score_validacao = 0.80
    elif campo == "local_nascimento" and parece_local_nascimento(valor):
        score_validacao = 0.75

    confianca = (
        0.40 * score_ocr +
        0.25 * score_rotulo +
        0.20 * peso_origem +
        0.15 * score_validacao
    )

    return round(min(max(confianca, 0.0), 1.0), 3)

# ============================================================
# OCR
# ============================================================

def executar_ocr(img, min_score: float = MIN_SCORE):
    resultado = ocr.predict(img)
    res = resultado[0]

    texts_raw = res.get("rec_texts", [])
    scores_raw = res.get("rec_scores", [])
    boxes_raw = res.get("rec_polys", res.get("dt_polys", []))

    filtrados = [
        (t, s, b)
        for t, s, b in zip(texts_raw, scores_raw, boxes_raw)
        if s >= min_score and str(t).strip()
    ]

    if not filtrados:
        return [], [], []

    texts, scores, boxes = zip(*filtrados)
    return list(texts), list(scores), list(boxes)


# ============================================================
# AGRUPAMENTO ESPACIAL EM LINHAS
# ============================================================

def centro_y(box):
    pts = np.array(box)
    return int(pts[:, 1].mean())


def centro_x(box):
    pts = np.array(box)
    return int(pts[:, 0].mean())


def altura_box(box):
    pts = np.array(box)
    return int(pts[:, 1].max() - pts[:, 1].min())


def agrupar_por_linha(texts, scores, boxes, tolerancia_y: int | None = None):
    if not texts:
        return []

    if tolerancia_y is None:
        alturas = [altura_box(b) for b in boxes]
        tolerancia_y = max(8, int(np.mean(alturas) * 0.65))

    tokens = sorted(
        zip(texts, scores, boxes),
        key=lambda tb: (centro_y(tb[2]), centro_x(tb[2]))
    )

    linhas = []
    atual = [tokens[0]]

    for token in tokens[1:]:
        cy_token = centro_y(token[2])
        cy_linha = int(np.mean([centro_y(t[2]) for t in atual]))

        if abs(cy_token - cy_linha) <= tolerancia_y:
            atual.append(token)
        else:
            linhas.append(sorted(atual, key=lambda tb: centro_x(tb[2])))
            atual = [token]

    linhas.append(sorted(atual, key=lambda tb: centro_x(tb[2])))
    return linhas


def montar_linhas_ocr(linhas_tokens):
    linhas = []

    for i, linha in enumerate(linhas_tokens):
        texto = " ".join(t for t, _, _ in linha)
        norm = aplicar_correcoes_ocr(texto)

        scores_linha = [float(s) for _, s, _ in linha]
        score_medio = float(np.mean([s for _, s, _ in linha]))

        tokens_linha = [
            {
                "texto": t,
                "score": float(s),
            }
            for t, s, _ in linha
        ]

        pontos = []
        for _, _, box in linha:
            pontos.extend(np.array(box).tolist())

        pts = np.array(pontos)

        linhas.append(
            LinhaOCR(
                indice=i,
                texto=texto,
                norm=norm,
                score_medio=score_medio,
                tokens=tokens_linha,
                x_min=int(pts[:, 0].min()),
                y_min=int(pts[:, 1].min()),
                x_max=int(pts[:, 0].max()),
                y_max=int(pts[:, 1].max()),
                cx=int(pts[:, 0].mean()),
                cy=int(pts[:, 1].mean()),
            )
        )

    return linhas


# ============================================================
# DETECÇÃO DE RÓTULOS
# ============================================================

def encontrar_rotulos_na_linha(linha_norm: str):
    encontrados = []

    for campo, aliases in ROTULOS_CAMPOS.items():
        for alias in aliases:
            alias_norm = aplicar_correcoes_ocr(alias)

            padrao = r"(?<![A-Z0-9])" + re.escape(alias_norm) + r"(?![A-Z0-9])"
            m = re.search(padrao, linha_norm)

            if m:
                encontrados.append({
                    "campo": campo,
                    "alias": alias_norm,
                    "pos_ini": m.start(),
                    "pos_fim": m.end(),
                    "score": 1.0,
                })
                continue

            sim = similaridade(linha_norm, alias_norm)

            if sim >= 0.90:
                encontrados.append({
                    "campo": campo,
                    "alias": alias_norm,
                    "pos_ini": 0,
                    "pos_fim": len(linha_norm),
                    "score": sim,
                })

    melhores = {}

    for item in encontrados:
        campo = item["campo"]
        atual = melhores.get(campo)

        if atual is None:
            melhores[campo] = item
        elif item["score"] > atual["score"]:
            melhores[campo] = item
        elif item["score"] == atual["score"] and item["pos_ini"] < atual["pos_ini"]:
            melhores[campo] = item

    return sorted(melhores.values(), key=lambda x: x["pos_ini"])


def trecho_do_rotulo_ate_proximo_rotulo(
    linha_norm: str,
    rotulo_atual: dict,
    rotulos_linha: list[dict],
) -> str:
    inicio = rotulo_atual["pos_fim"]

    proximos = [
        r["pos_ini"]
        for r in rotulos_linha
        if r["pos_ini"] > rotulo_atual["pos_fim"]
    ]

    fim = min(proximos) if proximos else len(linha_norm)

    trecho = linha_norm[inicio:fim]
    trecho = trecho.strip(" :.-|")

    return trecho


# ============================================================
# EXTRAÇÃO DE VALORES COM CONTEXTO
# ============================================================

def extrair_valor_mesma_linha(
    linha_norm: str,
    campo: str,
    rotulo_atual: dict,
    rotulos_linha: list[dict],
):
    """
    Extrai somente o trecho entre o rótulo atual e o próximo rótulo.
    Isso evita misturar:
      REGISTRO GERAL 28.XXX.102-5 2 VIA DATA DE EXPEDICAO 19/04/2024
    """
    trecho = trecho_do_rotulo_ate_proximo_rotulo(
        linha_norm=linha_norm,
        rotulo_atual=rotulo_atual,
        rotulos_linha=rotulos_linha,
    )

    if not trecho:
        return None

    return limpar_valor_por_campo(campo, trecho)


def extrair_valor_linha_abaixo(linhas: list[LinhaOCR], indice: int, campo: str):
    limite = min(indice + 4, len(linhas))

    for j in range(indice + 1, limite):
        proxima = linhas[j].norm.strip()

        if not proxima:
            continue

        rotulos_proxima = encontrar_rotulos_na_linha(proxima)
        if rotulos_proxima:
            break

        valor = limpar_valor_por_campo(campo, proxima)

        if valor:
            return valor, j

    return None, None


def extrair_campos_por_contexto(linhas: list[LinhaOCR]) -> dict:
    campos = {
        "nome": None,
        "cpf": None,
        "rg": None,
        "data_nascimento": None,
        "data_expedicao": None,
        "local_nascimento": None,
        "orgao_emissor": None,
        #"filiacao": None,
    }

    confianca = {}
    evidencias = {}
    linhas_consumidas = set()

    # 1. Associação forte por rótulo:
    #    mesma linha com recorte por próximo rótulo, depois linha abaixo.
    for i, linha in enumerate(linhas):
        rotulos = encontrar_rotulos_na_linha(linha.norm)

        if not rotulos:
            continue

        for rotulo in rotulos:
            campo = rotulo["campo"]

            if campo not in campos:
                continue

            if campos[campo]:
                continue

            valor = extrair_valor_mesma_linha(
                linha_norm=linha.norm,
                campo=campo,
                rotulo_atual=rotulo,
                rotulos_linha=rotulos,
            )

            origem = "mesma_linha"
            indice_valor = i

            if not valor:
                valor, indice_valor = extrair_valor_linha_abaixo(
                    linhas=linhas,
                    indice=i,
                    campo=campo,
                )
                origem = "linha_abaixo"

            if valor:
                campos[campo] = valor
                confianca[campo] = calcular_confianca_etiqueta(
                    campo=campo,
                    valor=valor,
                    score_ocr=linhas[indice_valor].score_medio,
                    score_rotulo=rotulo["score"],
                    origem=origem,
                )
                evidencias[campo] = {
                    "rotulo": rotulo["alias"],
                    "linha_rotulo": linha.texto,
                    "linha_valor": linhas[indice_valor].texto,
                    "indice_linha_valor": indice_valor,
                    "origem": origem,
                }
                linhas_consumidas.add(indice_valor)

    # 2. CPF por regex global.
    if not campos["cpf"]:
        texto_total = " ".join(l.norm for l in linhas)
        m = REGEX_CPF.search(texto_total)

        if m:
            cpf = limpar_cpf(m.group())

            if cpf:
                campos["cpf"] = cpf
                confianca["cpf"] = calcular_confianca_etiqueta(
                    campo="cpf",
                    valor=cpf,
                    score_ocr=0.80,
                    score_rotulo=0.0,
                    origem="regex_global",
                )
                evidencias["cpf"] = {
                    "valor_bruto": m.group(),
                    "linha": linha.texto,
                    "indice_linha_valor": i,
                    "origem": "regex_global",
                }

    # 3. RG por fallback, somente se ainda não veio do rótulo.
#    if not campos["rg"]:
#        cpf_digitos = re.sub(r"\D", "", campos["cpf"] or "")
#
#        for i, linha in enumerate(linhas):
#            if i in linhas_consumidas:
#                continue
#
#            m = REGEX_RG.search(linha.norm)
#            if not m:
#                continue
#
#            rg = limpar_rg(m.group())
#
#            if rg and rg != cpf_digitos:
#                campos["rg"] = rg
#                confianca["rg"] = calcular_confianca_etiqueta(
#                    campo="rg",
#                    valor=rg,
#                    score_ocr=linha.score_medio,
#                    score_rotulo=0.0,
#                    origem="regex_linha",
#                )
#                evidencias["rg"] = {
#                    "linha": linha.texto,
#                    "origem": "regex_linha",
#                }
#                break

#    # 4. Datas soltas
#    datas_soltas = []
#
#    for i, linha in enumerate(linhas):
#        if i in linhas_consumidas:
#            continue
#
#        if encontrar_rotulos_na_linha(linha.norm):
#            continue
#
#        for m in REGEX_DATA.finditer(linha.norm):
#            data = limpar_data(m.group())
#
#            if data:
#                datas_soltas.append({
#                    "data": data,
#                    "linha": linha.texto,
#                    "indice": i,
#                })

    # Fallback: se for a única data solta, será atribuida para nascimento 
#    if not campos["data_nascimento"] and len(datas_soltas) == 1:
#        campos["data_nascimento"] = datas_soltas[0]["data"]
#        confianca["data_nascimento"] = calcular_confianca_etiqueta(
#            campo="data_nascimento",
#            valor=datas_soltas[0]["data"],
#            score_ocr=linhas[datas_soltas[0]["indice"]].score_medio,
#            score_rotulo=0.0,
#            origem="fallback_unica_data_solta",
#        )
#        evidencias["data_nascimento"] = {
#            "linha": datas_soltas[0]["linha"],
#            "origem": "fallback_unica_data_solta",
#        }

    # 5. Local de nascimento por padrão municipal/UF, apenas fallback.
    if not campos["local_nascimento"]:
        for linha in linhas:
            if REGEX_LOCAL_UF.search(linha.norm):
                valor = limpar_texto(linha.norm)

                if parece_local_nascimento(valor):
                    campos["local_nascimento"] = valor
                    confianca["local_nascimento"] =  calcular_confianca_etiqueta(
                        campo="local_nascimento",
                        valor=valor,
                        score_ocr=linha.score_medio,
                        score_rotulo=0.0,
                        origem="regex_local_uf",
                    )
                    evidencias["local_nascimento"] = {
                    "linha": linha.texto,
                    "indice_linha_valor": i,
                    "origem": "regex_local_uf",
                    }
                    break

    return {
        "campos": campos,
        "confianca": confianca,
        "evidencias": evidencias,
        "linhas": [asdict(l) for l in linhas],
    }


# ============================================================
# CLASSIFICAÇÃO DO TIPO DE DOCUMENTO
# ============================================================

MARCADORES_DOCUMENTO = {
    "RG": [
        "CARTEIRA DE IDENTIDADE",
        "CEDULA DE IDENTIDADE",
        "REGISTRO GERAL",
        "SECRETARIA DE SEGURANCA PUBLICA",
        "INSTITUTO DE IDENTIFICACAO",
    ],
    "CNH": [
        "CARTEIRA NACIONAL DE HABILITACAO",
        "REGISTRO NACIONAL DE CARTEIRA DE HABILITACAO",
        "DETRAN",
        "SENATRAN",
        "DENATRAN",
        "HABILITACAO",
    ],
}


def classificar_tipo_documento(texts, scores, min_score: float = MIN_SCORE):
    votos = {"RG": 0.0, "CNH": 0.0}
    evidencias = []

    for text, score in zip(texts, scores):
        if score < min_score:
            continue

        texto_norm = aplicar_correcoes_ocr(text)

        for tipo, marcadores in MARCADORES_DOCUMENTO.items():
            for marcador in marcadores:
                marcador_norm = aplicar_correcoes_ocr(marcador)
                sim = similaridade(texto_norm, marcador_norm)

                if sim >= 0.75 or marcador_norm in texto_norm:
                    votos[tipo] += sim
                    evidencias.append({
                        "tipo": tipo,
                        "texto": text,
                        "marcador": marcador,
                        "similaridade": round(sim, 3),
                    })

    total_votos = sum(votos.values())

    if total_votos == 0:
        return {
            "tipo": "DESCONHECIDO",
            "confianca": 0.0,
            "votos": votos,
            "votos_normalizados": {"RG": 0.0, "CNH": 0.0},
            "evidencias": evidencias,
        }

    votos_normalizados = {
        tipo: round(valor / total_votos, 3)
        for tipo, valor in votos.items()
    }

    tipo_final = max(votos_normalizados, key=votos_normalizados.get)

    return {
        "tipo": tipo_final,
        "confianca": votos_normalizados[tipo_final],
        "votos": votos,
        "votos_normalizados": votos_normalizados,
        "evidencias": evidencias,
    }

# ============================================================
# XML
# ============================================================

def resultado_para_xml(resultado: dict) -> str:
    root = ET.Element("documento")

    arquivo_el = ET.SubElement(root, "arquivo")
    arquivo_el.text = resultado.get("arquivo", "")

    tipo_el = ET.SubElement(root, "tipo_documento")
    tipo_el.text = resultado["tipo_documento"]["tipo"]

    campos_el = ET.SubElement(root, "campos")

    campos = resultado["extracao"]["campos"]
    confianca = resultado["extracao"]["confianca"]
    evidencias = resultado["extracao"]["evidencias"]

    for nome_campo, valor in campos.items():
        campo_el = ET.SubElement(campos_el, nome_campo)
        campo_el.text = valor or ""

        if nome_campo in confianca:
            campo_el.set("confianca", str(confianca[nome_campo]))

        ev = evidencias.get(nome_campo)

        if isinstance(ev, dict):
            campo_el.set("origem", ev.get("origem", ""))

            if ev.get("rotulo"):
                campo_el.set("rotulo", ev["rotulo"])

            if ev.get("linha_rotulo"):
                campo_el.set("linha_rotulo", ev["linha_rotulo"])

            if ev.get("linha_valor"):
                campo_el.set("linha_valor", ev["linha_valor"])

            if ev.get("linha"):
                campo_el.set("linha", ev["linha"])

    linhas_el = ET.SubElement(root, "linhas_ocr")

    for linha in resultado["extracao"]["linhas"]:
        linha_el = ET.SubElement(linhas_el, "linha")
        linha_el.set("indice", str(linha["indice"]))

        if linha.get("x_min") is not None:
            linha_el.set("x_min", str(linha["x_min"]))
            linha_el.set("y_min", str(linha["y_min"]))
            linha_el.set("x_max", str(linha["x_max"]))
            linha_el.set("y_max", str(linha["y_max"]))

        linha_el.text = linha["texto"]

    return ET.tostring(root, encoding="unicode")


def salvar_xml(resultado: dict, caminho_xml: str):
    xml_str = resultado_para_xml(resultado)

    with open(caminho_xml, "w", encoding="utf-8") as f:
        f.write(xml_str)

    return caminho_xml


# ============================================================
# PROCESSAMENTO DE DOCUMENTO
# ============================================================

def processar_documento(caminho_imagem: str, min_score: float = MIN_SCORE):
    img = cv2.imread(caminho_imagem)

    if img is None:
        raise ValueError(f"Nao foi possivel carregar a imagem: {caminho_imagem}")

    texts, scores, boxes = executar_ocr(img, min_score=min_score)

    tipo_documento = classificar_tipo_documento(
        texts=texts,
        scores=scores,
        min_score=min_score,
    )

    linhas_tokens = agrupar_por_linha(texts, scores, boxes)
    linhas = montar_linhas_ocr(linhas_tokens)

    texto_por_linhas = [
        linha.texto.strip()
        for linha in linhas
        if linha.texto.strip()
    ]

    extracao = extrair_campos_por_contexto(linhas)

    return {
        "arquivo": os.path.basename(caminho_imagem),
        "tipo_documento": tipo_documento,
        "extracao": extracao,
        "texto_por_linhas": texto_por_linhas,
        "tokens_ocr": [
            {
                "texto": t,
                "score": float(s),
            }
            for t, s in zip(texts, scores)
        ],
    }


# ============================================================
# PROCESSAMENTO EM LOTE COM LOG
# ============================================================

from pathlib import Path
from datetime import datetime
import os
from openpyxl import Workbook


def valor_planilha(valor):
    if valor is None:
        return "-"
    if isinstance(valor, str) and valor.strip().lower() == "none":
        return "-"
    if valor == "":
        return "-"
    return valor


def extrair_confianca_ocr(campo, evidencias, linhas):
    """
    Retorna a confiança OCR média dos tokens da linha usada
    como evidência para o campo.
    """

    ev = evidencias.get(campo)

    if not isinstance(ev, dict):
        return "-"

    indice_linha = ev.get("indice_linha_valor")

    if indice_linha is None:
        return "-"

    for linha in linhas:
        if linha.get("indice") == indice_linha:
            tokens = linha.get("tokens", [])

            scores = [
                float(t.get("score"))
                for t in tokens
                if isinstance(t, dict) and t.get("score") is not None
            ]

            if scores:
                return round(sum(scores) / len(scores), 4)

    return "-"

def criar_layout_planilha(ws):
    campos_planilha = [
        "nome",
        "cpf",
        "rg",
        "data_nascimento",
        "local_nascimento",
        "orgao_emissor",
    ]

    ws.cell(row=1, column=1, value="campo / id documento")

    ws.cell(row=2, column=1, value="tipo detectado")
    ws.cell(row=3, column=1, value="votos_normalizados")

    mapa_linhas = {
        "tipo_detectado": 2,
        "votos_normalizados": 3,
    }

    linha_atual = 4

    for campo in campos_planilha:
        ws.cell(row=linha_atual, column=1, value=campo)
        ws.cell(row=linha_atual, column=2, value="confianca_ocr")
        mapa_linhas[(campo, "confianca_ocr")] = linha_atual

        linha_atual += 1

        ws.cell(row=linha_atual, column=2, value="confianca_campo")
        mapa_linhas[(campo, "confianca_campo")] = linha_atual

        linha_atual += 1

    return mapa_linhas


def preencher_coluna_planilha(
    ws,
    coluna_documento,
    id_documento,
    tipo_doc,
    campos,
    confianca,
    evidencias,
    linhas,
    mapa_linhas,
):
    ws.cell(row=1, column=coluna_documento, value=id_documento)

    ws.cell(
        row=mapa_linhas["tipo_detectado"],
        column=coluna_documento,
        value=valor_planilha(tipo_doc.get("tipo")),
    )

    ws.cell(
        row=mapa_linhas["votos_normalizados"],
        column=coluna_documento,
        value=valor_planilha(str(tipo_doc.get("votos_normalizados"))),
    )

    for campo in campos.keys():
        if (campo, "confianca_ocr") not in mapa_linhas:
            continue

        confianca_ocr = extrair_confianca_ocr(
            campo=campo,
            evidencias=evidencias,
            linhas=linhas,
        )

        confianca_campo = confianca.get(campo)

        ws.cell(
            row=mapa_linhas[(campo, "confianca_ocr")],
            column=coluna_documento,
            value=valor_planilha(confianca_ocr),
        )

        ws.cell(
            row=mapa_linhas[(campo, "confianca_campo")],
            column=coluna_documento,
            value=valor_planilha(confianca_campo),
        )

def processar_pasta(
    pasta_entrada: str,
    pasta_saida: str = "saida",
    min_score: float = MIN_SCORE,
    nome_log: str | None = None,
    nome_planilha: str | None = None,
):
    os.makedirs(pasta_saida, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if nome_log is None:
        nome_log = f"resultado_ocr_{timestamp}.log"

    if nome_planilha is None:
        nome_planilha = f"resultado_ocr_{timestamp}.xlsx"

    caminho_log = os.path.join(pasta_saida, nome_log)
    caminho_planilha = os.path.join(pasta_saida, nome_planilha)

    arquivos = sorted([
        p for p in Path(pasta_entrada).iterdir()
        if p.is_file() and p.suffix.lower() in EXTENSOES_IMAGEM
    ])

    resultados = []

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados OCR"

    mapa_linhas_planilha = criar_layout_planilha(ws)

    with open(caminho_log, "w", encoding="utf-8") as log:
        log.write("=" * 80 + "\n")
        log.write("LOG DE PROCESSAMENTO OCR\n")
        log.write(f"Data/hora: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        log.write(f"Pasta entrada: {pasta_entrada}\n")
        log.write(f"Total de arquivos encontrados: {len(arquivos)}\n")
        log.write(f"Min score: {min_score}\n")
        log.write("=" * 80 + "\n\n")

        if not arquivos:
            log.write("Nenhuma imagem encontrada.\n")
            wb.save(caminho_planilha)
            return resultados

        for idx, caminho in enumerate(arquivos, 1):
            log.write("\n" + "=" * 80 + "\n")
            log.write(f"[{idx}/{len(arquivos)}] ARQUIVO: {caminho.name}\n")
            log.write("=" * 80 + "\n")

            try:
                resultado = processar_documento(
                    caminho_imagem=str(caminho),
                    min_score=min_score,
                )

                resultados.append(resultado)

                tipo_doc = resultado.get("tipo_documento", {})
                extracao = resultado.get("extracao", {})
                campos = extracao.get("campos", {})
                confianca = extracao.get("confianca", {})
                evidencias = extracao.get("evidencias", {})
                linhas = extracao.get("linhas", [])
                tokens = resultado.get("tokens_ocr", [])

                log.write("\nTIPO_DOCUMENTO\n")
                log.write(f"  tipo: {tipo_doc.get('tipo')}\n")
                log.write(f"  votos_normalizados: {tipo_doc.get('votos_normalizados')}\n")

                log.write("\nCAMPOS_EXTRAIDOS\n")
                for campo, valor in campos.items():
                    log.write(f"  {campo}: {valor}\n")

                log.write("\nCONFIANCA\n")
                for campo, valor in confianca.items():
                    log.write(f"  {campo}: {valor}\n")

                log.write("\nEVIDENCIAS\n")
                for campo, ev in evidencias.items():
                    log.write(f"  {campo}:\n")
                    if isinstance(ev, dict):
                        for k, v in ev.items():
                            log.write(f"    {k}: {v}\n")
                    else:
                        log.write(f"    {ev}\n")

                log.write("\nLINHAS_OCR_AGRUPADAS_NORMALIZADAS:\n")
                for linha in linhas:
                    log.write(f"{linha.get('norm')}\n")

                log.write("\nTOKENS_OCR\n")
                for token in tokens:
                    score = token.get("score")
                    texto = token.get("texto")

                    if score is not None:
                        log.write(f"  [{score:.2f}] {texto}\n")
                    else:
                        log.write(f"  [-] {texto}\n")

                coluna_documento = idx + 2

                preencher_coluna_planilha(
                    ws=ws,
                    coluna_documento=coluna_documento,
                    id_documento=idx,
                    tipo_doc=tipo_doc,
                    campos=campos,
                    confianca=confianca,
                    evidencias=evidencias,
                    linhas=linhas,
                    mapa_linhas=mapa_linhas_planilha,
                )

                log.write("\nSTATUS: OK\n")
                print(f"Processado: {caminho.name}")

            except Exception as e:
                log.write("\nSTATUS: ERRO\n")
                log.write(f"Erro ao processar {caminho.name}: {e}\n")
                print(f"Erro ao processar {caminho.name}: {e}")

    wb.save(caminho_planilha)

    print(f"Log salvo em: {caminho_log}")
    print(f"Planilha salva em: {caminho_planilha}")

    return resultados




# ============================================================
# TESTE RÁPIDO SEM OCR
# ============================================================

def testar_com_linhas_mock():
    linhas_texto = [
        "SECRETARIA DA SEGURANÇA PÚBLICA REGISTRO GERAL 28.507.102-5 2 VIA DATADE EXPEDIÇÃO 19/04/2024",
        "NOME MARIA DA SILVA",
        "DATA DE NASCIMENTO",
        "10/05/1990",
        "CPF 123.456.789-00",
        "NATURALIDADE SAO PAULO SP",
    ]

    linhas = []

    for i, texto in enumerate(linhas_texto):
        linhas.append(
            LinhaOCR(
                indice=i,
                texto=texto,
                norm=aplicar_correcoes_ocr(texto),
                x_min=0,
                y_min=i * 30,
                x_max=900,
                y_max=i * 30 + 20,
                cx=450,
                cy=i * 30 + 10,
            )
        )

    extracao = extrair_campos_por_contexto(linhas)

    resultado = {
        "arquivo": "teste_mock.jpg",
        "tipo_documento": {
            "tipo": "RG",
            "votos": {},
            "evidencias": [],
        },
        "extracao": extracao,
        "tokens_ocr": [],
    }

    print(json.dumps(extracao["campos"], ensure_ascii=False, indent=2))
    print(resultado_para_xml(resultado))


if __name__ == "__main__":
    testar_com_linhas_mock()

    # Para processar imagens reais:

    resultados = processar_pasta(
         pasta_entrada="privado/1901_ate_2161",
         pasta_saida="privado",
         min_score=0.50,
         nome_log="21Mai_log_codex_1901_ate_2161_v1.log",
     )